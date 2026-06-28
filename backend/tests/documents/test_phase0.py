"""Phase 0 acceptance tests — see docs/documents/00-document-upload-plan.md.

No pytest-asyncio in this repo, so async coroutines are driven with
asyncio.run() from plain sync test functions.
"""

from __future__ import annotations

import asyncio
import io
from typing import Any

import pytest

from app.documents.base import UnsupportedDocumentError
from app.documents.blob_store import DiskBlobStore, S3BlobStore
from app.documents.processors.local import LocalProcessor
from app.documents.service import DocumentService, ThreadOwnershipError
from app.documents.store import DocumentNotFoundError, InMemoryDocumentStore
from app.services.thread_store import InMemoryThreadStore


def test_disk_blob_store_round_trip(tmp_path):
    store = DiskBlobStore(tmp_path)

    async def scenario():
        uri = await store.put(key="u1/doc1", data=b"hello", content_type="text/plain")
        assert uri.startswith("file://")
        assert await store.get(key="u1/doc1") == b"hello"
        await store.delete(key="u1/doc1")
        with pytest.raises(FileNotFoundError):
            await store.get(key="u1/doc1")

    asyncio.run(scenario())


def test_disk_blob_store_rejects_traversal(tmp_path):
    store = DiskBlobStore(tmp_path)
    with pytest.raises(ValueError):
        asyncio.run(store.put(key="../escape", data=b"x", content_type="text/plain"))


class _FakeBody:
    def __init__(self, data: bytes) -> None:
        self._data = data
        self.closed = False

    def read(self) -> bytes:
        return self._data

    def close(self) -> None:
        self.closed = True


class _FakeS3Client:
    def __init__(self) -> None:
        self.objects: dict[tuple[str, str], tuple[bytes, str]] = {}

    def put_object(self, **kwargs: Any) -> None:
        self.objects[(kwargs["Bucket"], kwargs["Key"])] = (
            kwargs["Body"],
            kwargs["ContentType"],
        )

    def get_object(self, **kwargs: Any) -> dict[str, Any]:
        data, _content_type = self.objects[(kwargs["Bucket"], kwargs["Key"])]
        return {"Body": _FakeBody(data)}

    def delete_object(self, **kwargs: Any) -> None:
        self.objects.pop((kwargs["Bucket"], kwargs["Key"]), None)

    def generate_presigned_url(self, operation: str, **kwargs: Any) -> str:
        assert operation == "get_object"
        params = kwargs["Params"]
        return f"https://signed.test/{params['Bucket']}/{params['Key']}?ttl={kwargs['ExpiresIn']}"


def test_s3_blob_store_round_trip():
    client = _FakeS3Client()
    store = S3BlobStore(
        bucket="bucket",
        endpoint_url="https://storage.example",
        region_name="auto",
        access_key_id="key",
        secret_access_key="secret",
        client=client,
    )

    async def scenario():
        uri = await store.put(key="u1/doc1", data=b"hello", content_type="text/plain")
        assert uri == "s3://bucket/u1/doc1"
        assert await store.get(key="u1/doc1") == b"hello"
        assert await store.signed_url(key="u1/doc1", ttl_s=600) == (
            "https://signed.test/bucket/u1/doc1?ttl=600"
        )
        await store.delete(key="u1/doc1")
        assert ("bucket", "u1/doc1") not in client.objects

    asyncio.run(scenario())


def test_local_processor_parses_text():
    proc = LocalProcessor()
    result = asyncio.run(
        proc.parse(data=b"# Title\n\nBody text.", filename="note.md", mime_type="text/markdown")
    )
    assert "Body text." in result.markdown
    assert result.summary
    assert result.provider == "local"


def test_local_processor_parses_docx():
    from docx import Document as DocxDocument

    buf = io.BytesIO()
    doc = DocxDocument()
    doc.add_paragraph("Jane Mentee — Software Engineer")
    doc.add_paragraph("Skills: Python, FastAPI")
    doc.save(buf)

    proc = LocalProcessor()
    result = asyncio.run(
        proc.parse(
            data=buf.getvalue(),
            filename="cv.docx",
            mime_type=(
                "application/vnd.openxmlformats-officedocument"
                ".wordprocessingml.document"
            ),
        )
    )
    assert "Jane Mentee" in result.markdown
    assert "Python" in result.markdown


def test_local_processor_rejects_images():
    # Images need vision OCR (OpenAIDocumentProcessor); the local parser can't.
    proc = LocalProcessor()
    assert not proc.supports("image/png")
    with pytest.raises(UnsupportedDocumentError):
        asyncio.run(
            proc.parse(data=b"\x89PNG", filename="x.png", mime_type="image/png")
        )


def test_in_memory_store_is_user_scoped():
    store = InMemoryDocumentStore()

    async def scenario():
        doc = await store.create_document(
            user_id="user-a",
            thread_id="thread-1",
            purpose="chat_attachment",
            filename="a.pdf",
            mime_type="application/pdf",
            size_bytes=10,
            file_hash="abc",
        )
        # Owner can read; another user cannot.
        assert (await store.get_document(doc.id, "user-a")).id == doc.id
        with pytest.raises(DocumentNotFoundError):
            await store.get_document(doc.id, "user-b")

    asyncio.run(scenario())


def test_thread_ownership_guard():
    threads = InMemoryThreadStore()
    service = DocumentService(
        store=InMemoryDocumentStore(),
        blobs=DiskBlobStore("/tmp/does-not-matter"),
        processor=LocalProcessor(),
        threads=threads,
    )

    async def scenario():
        thread = await threads.create_thread("owner")
        # Owner passes; non-owner is rejected (plan decision #15).
        await service.assert_thread_owned(user_id="owner", thread_id=thread.id)
        with pytest.raises(ThreadOwnershipError):
            await service.assert_thread_owned(user_id="intruder", thread_id=thread.id)

    asyncio.run(scenario())


def test_profile_context_injects_cv_markdown_and_about_me():
    store = InMemoryDocumentStore()
    service = DocumentService(
        store=store,
        blobs=DiskBlobStore("/tmp/does-not-matter"),
        processor=LocalProcessor(),
        threads=InMemoryThreadStore(),
    )

    async def scenario():
        # No profile → nothing injected.
        assert await service.build_profile_context(user_id="u1") == ""
        assert await service.build_about_context(user_id="u1") == ""

        # A ready CV document carrying its Markdown transcription, activated.
        doc = await store.create_document(
            user_id="u1", thread_id=None, purpose="profile_cv",
            filename="cv.pdf", mime_type="application/pdf",
            size_bytes=10, file_hash="h",
        )
        await store.update_document(
            doc.id, status="ready", extracted_text="# Jane\nSoftware Engineer"
        )
        await store.set_cv(user_id="u1", cv_document_id=doc.id)
        ctx = await service.build_profile_context(user_id="u1")
        assert "Jane" in ctx and "Software Engineer" in ctx

        # Free-text about-me is injected once set.
        await service.set_about_me(user_id="u1", about_me="First-gen CS student.")
        assert "First-gen CS student." in await service.build_about_context(user_id="u1")

    asyncio.run(scenario())
